"""有限额、无外部链接的 XLSX / JSON / ZIP 导入预检。不会写入正品记录。"""
import sys, json, zipfile, io, stat, unicodedata, re
from pathlib import PurePosixPath, Path
from asset_worker import asset, resource_check
MAX_UNPACKED=500*1024**2
IMAGE_SUFFIXES=('.jpg','.jpeg','.png','.webp')

def filename_fields(name):
    stem=PurePosixPath(name).stem.strip()
    parts=re.split(r'\s+',stem)
    if len(parts) not in (2,3) or not all(parts):
        raise ValueError('图片文件名须为“物料编号 型号 [供应商]”两段或三段，各字段内部不要使用空格。')
    if any(any(ord(c)<32 for c in s) for s in parts):
        raise ValueError('文件名不能包含控制字符。')
    return dict(materialCode=parts[0],modelNumber=parts[1],supplierName=parts[2] if len(parts)==3 else '')

def zip_entries(data):
    z=zipfile.ZipFile(io.BytesIO(data)); entries={}; total=0
    if len(z.infolist())>5000: raise ValueError('压缩包条目超过 5000 个，请拆分导入。')
    for info in z.infolist():
        name=unicodedata.normalize('NFC',info.filename)
        if '\\' in name or name.startswith('/') or ':' in name or '..' in PurePosixPath(name).parts or stat.S_ISLNK(info.external_attr>>16) or info.flag_bits&1:
            raise ValueError('压缩包含不安全路径、符号链接或加密文件。')
        if info.is_dir() or name.startswith('__MACOSX/') or PurePosixPath(name).name.startswith('.'): continue
        if name in entries: raise ValueError('压缩包中存在重复文件路径。')
        total+=info.file_size
        if total>MAX_UNPACKED or info.file_size>100*1024**2: raise ValueError('解压后数据过大，请分批导入。')
        entries[name]=info
    return z, entries

def manifest(data,name):
    if name.lower().endswith('.json'):
        parsed=json.loads(data.decode('utf-8-sig'))
        if not isinstance(parsed,dict) or parsed.get('schemaVersion')!='1.0' or not isinstance(parsed.get('records'),list): raise ValueError('JSON 须包含 schemaVersion: "1.0" 和 records 数组。')
        return parsed['records']
    zip_entries(data) # 读取工作簿前限制其内部 ZIP 体积。
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        ws=wb['标准清单'] if '标准清单' in wb.sheetnames else wb.worksheets[0]
        if ws.max_row and ws.max_row>2001: raise ValueError('每次最多 2000 条标准，请拆分工作簿。')
        rows=ws.iter_rows(); headers=[str(c.value or '').strip() for c in next(rows)]
        aliases={'物料编号':'materialCode','供应商名称':'supplierName','型号':'modelNumber','正品图片路径':'imagePath','检测说明':'description','物料名称':'name','正品丝印文字':'markingText'}
        if not all(h in headers for h in ('物料编号','型号','正品图片路径')): raise ValueError('Excel 缺少必填列，请下载模板。')
        result=[]
        for cells in rows:
            if len(result)>2000: raise ValueError('每次最多 2000 条标准。')
            if not any(c.value is not None for c in cells): continue
            row={}; error=''
            for i,c in enumerate(cells):
                if i>=len(headers) or headers[i] not in aliases: continue
                key=aliases[headers[i]]
                if c.data_type=='f': error='不支持公式，请粘贴为文本值。'
                if key in ('materialCode','modelNumber') and c.value is not None and not isinstance(c.value,str): error='物料编号和型号必须为文本，避免前导零丢失。请将单元格设为文本后重新输入。'
                row[key]=str(c.value or '').strip()
            row['_error']=error;result.append(row)
        return result
    finally: wb.close()

def main(r):
    files=r['files']; contents={f['name']:Path(f['path']).read_bytes() for f in files}
    filename_mode=False
    if len(contents)!=len(files): raise ValueError('上传文件名重复，请打包为 ZIP 并使用不同路径。')
    if r.get('archiveFiles'):
        manifests=[n for n in contents if n.lower().endswith(('.xlsx','.json'))]
        if len(manifests)>1: raise ValueError('RAR 中最多放一个 XLSX 或 JSON 清单。')
        if manifests: rows=manifest(contents[manifests[0]],manifests[0])
        else:
            filename_mode=True
            rows=[{'_filename':n,'imagePath':n} for n in sorted(contents) if n.lower().endswith(IMAGE_SUFFIXES)]
        def picture(name):
            candidates=[name,str(PurePosixPath(manifests[0]).parent/name)] if manifests else [name]
            for key in candidates:
                if key in contents:return contents[key]
            raise ValueError('找不到图片：'+name)
    elif len(files)==1 and files[0]['name'].lower().endswith('.zip'):
        archive,entries=zip_entries(contents[files[0]['name']]); manifests=[n for n in entries if n.lower().endswith(('.xlsx','.json'))]
        if len(manifests)>1: raise ValueError('ZIP 中最多放一个 XLSX 或 JSON 清单，或只放按名称整理的图片。')
        if manifests:
            mf=manifests[0];data=archive.read(entries[mf]); rows=manifest(data,mf)
        else:
            filename_mode=True;mf=''
            rows=[{'_filename':n,'imagePath':n} for n in sorted(entries) if n.lower().endswith(IMAGE_SUFFIXES)]
        def picture(name):
            candidates=[name,str(PurePosixPath(mf).parent/name)]
            for key in candidates:
                if key in entries: return archive.read(entries[key])
            raise ValueError('找不到图片：'+name)
    else:
        manifests=[n for n in contents if n.lower().endswith(('.xlsx','.json'))]
        if len(manifests)!=1: raise ValueError('请选择一个 Excel 或 JSON 清单，并同时选取清单中使用的图片；也可上传完整 ZIP。')
        rows=manifest(contents[manifests[0]],manifests[0])
        def picture(name):
            if name in contents: return contents[name]
            matches=[v for k,v in contents.items() if PurePosixPath(k).name==PurePosixPath(name).name]
            if len(matches)!=1: raise ValueError('图片缺失或文件名重复：'+name)
            return matches[0]
    if not rows or len(rows)>2000: raise ValueError('清单不能为空，单次最多 2000 条。建议每批 500 条。')
    output=[];assets={};cache={}
    for i,raw in enumerate(rows):
        if i%20==0: resource_check()
        row={'row':i+(1 if filename_mode else 2),'material':'','supplier':'','modelNumber':'','name':'','description':'','lines':[],'imagePath':raw.get('imagePath','') if isinstance(raw,dict) else '', 'status':'ready','message':''}
        try:
            if not isinstance(raw,dict): raise ValueError('记录必须为对象。')
            if '_filename' in raw:
                # 文件名错误也保留可解码图片，后续可在待处理详情补齐资料。
                imagePath=raw['imagePath']
                if imagePath not in cache: cache[imagePath]=asset(picture(imagePath),imagePath,r['folder'])
                a=cache[imagePath];assets[a['id']]=a;row['assetId']=a['id']
                raw={**filename_fields(raw['_filename']),'imagePath':imagePath}
            if raw.get('_error'): raise ValueError(raw['_error'])
            unknown=set(raw)-{'materialCode','supplierName','modelNumber','imagePath','description','name','markingText','markingLines','_error'}
            if unknown: raise ValueError('不支持的字段：'+', '.join(sorted(unknown)))
            for source,target,maxlen in [('materialCode','material',160),('supplierName','supplier',160),('modelNumber','modelNumber',200),('name','name',200),('description','description',8000),('imagePath','imagePath',500)]:
                value=raw.get(source,'')
                if not isinstance(value,str): raise ValueError(source+' 必须为文本。')
                row[target]=unicodedata.normalize('NFC',value.strip())
                if len(value)>maxlen: raise ValueError(source+' 过长。')
            if not row['material'] or not row['modelNumber'] or not row['imagePath']: raise ValueError('物料编号、型号、图片路径均必填。供应商选填。')
            imagePath=row['imagePath']
            if '\\' in imagePath or imagePath.startswith('/') or ':' in imagePath or '..' in PurePosixPath(imagePath).parts: raise ValueError('图片路径必须是清单内的安全相对路径。')
            lines=raw.get('markingLines',raw.get('markingText','').splitlines())
            if not isinstance(lines,list) or len(lines)>50 or any(not isinstance(x,str) or len(x)>300 for x in lines): raise ValueError('丝印文字格式不正确。')
            row['lines']=[x.strip() for x in lines if x.strip()]
            if imagePath not in cache: cache[imagePath]=asset(picture(imagePath),imagePath,r['folder'])
            a=cache[imagePath];assets[a['id']]=a;row['assetId']=a['id']
        except Exception as e: row['status']='error';row['message']=str(e)
        output.append(row)
    if len(files)==1 and files[0]['name'].lower().endswith('.zip'): archive.close()
    return {'rows':output,'assets':list(assets.values())}

if __name__=='__main__':
    sys.stdin.reconfigure(encoding='utf-8-sig')
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8', errors='backslashreplace')
    try: print(json.dumps(main(json.load(sys.stdin)),ensure_ascii=False))
    except Exception as e: print(json.dumps({'error':str(e)},ensure_ascii=False));sys.exit(1)
